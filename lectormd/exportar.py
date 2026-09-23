"""Exportación a Word (.docx).

El PDF lo imprime directamente el motor web (ver app.py). Aquí vive el
conversor a Word: recorre los tokens de markdown-it —no el HTML— y construye
el documento con python-docx.

Qué se traduce y cómo:
  - Encabezados, párrafos, negrita, cursiva, tachado, enlaces.
  - Listas con numeración real de Word (cada lista reinicia en 1).
  - Código con los colores de Pygments.
  - Tablas con cabecera y alineación por columna.
  - Citas y avisos [!NOTE] con borde lateral de color.
  - Fórmulas como ecuaciones nativas de Word (editables), vía MathML → OMML.
  - Diagramas Mermaid como imagen PNG que genera el propio visor.
  - Imágenes locales y remotas.
"""

from __future__ import annotations

import base64
import io
import re
import sys
import urllib.request
from pathlib import Path
from urllib.parse import unquote, urlparse

import latex2mathml.converter
import mathml2omml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as ImagenDocx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Pt, RGBColor
from pygments.lexers import get_lexer_by_name, guess_lexer
from pygments.styles import get_style_by_name
from pygments.util import ClassNotFound

from . import renderer

# ------------------------------------------------------------------ paleta
ACENTO = RGBColor(0x4A, 0x56, 0xC8)
TEXTO = RGBColor(0x1D, 0x21, 0x26)
SUAVE = RGBColor(0x5C, 0x65, 0x70)
FONDO_CODIGO = "F6F7F9"
FONDO_CODIGO_LINEA = "EEF0F3"
BORDE = "D8DCE2"
FUENTE = "Calibri"
# OOXML no tiene cadena de fuentes alternativas: si la fuente falta, Word la
# sustituye por otra de paso fijo, pero LibreOffice cae a una serif. Se usa la
# que seguro existe donde se exporta: Consolas viene con Office en Windows y
# Mac; Liberation Mono viene con LibreOffice en Linux.
FUENTE_MONO = "Consolas" if sys.platform in ("win32", "darwin") else "Liberation Mono"
PANOSE_MONO = {"Consolas": "020B0609020204030204", "Liberation Mono": "02070309020205020404"}

AVISOS = {
    "note": ("Nota", "3B82F6"),
    "tip": ("Consejo", "10B981"),
    "important": ("Importante", "8B5CF6"),
    "warning": ("Atención", "F59E0B"),
    "caution": ("Cuidado", "EF4444"),
}

# A4 con márgenes de 2,2 cm: ancho útil para imágenes y diagramas.
MARGEN = Cm(2.2)
ANCHO_UTIL = Cm(21.0) - 2 * MARGEN
EMU_POR_PX = 9525  # a 96 ppp

# Orden obligatorio de los hijos de <w:pPr> y <w:rPr>: Word rechaza el
# archivo ("contenido ilegible") si se insertan desordenados.
_ORDEN_PPR = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
    "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
]
_ORDEN_RPR = [
    "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
    "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss",
    "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden",
    "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
    "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
    "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
    "w:specVanish", "w:oMath",
]

_ORDEN_TBLPR = [
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
    "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
    "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
    "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
]
_ORDEN_TCPR = [
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
    "w:vAlign", "w:hideMark",
]

FORMATOS_DOCX = {"png", "jpeg", "gif", "bmp", "tiff", "x-wmf"}


# ------------------------------------------------------------ utilidades XML
def _insertar(padre, elemento, nombre: str, orden: list[str]) -> None:
    """Inserta respetando el orden del esquema, reemplazando si ya existe."""
    for viejo in padre.findall(qn(nombre)):
        padre.remove(viejo)
    sucesores = orden[orden.index(nombre) + 1 :]
    padre.insert_element_before(elemento, *sucesores)


def _elemento(nombre: str, **attrs) -> OxmlElement:
    el = OxmlElement(nombre)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


def _sombrear_parrafo(p, color: str) -> None:
    _insertar(p._p.get_or_add_pPr(), _elemento("w:shd", val="clear", color="auto", fill=color),
              "w:shd", _ORDEN_PPR)


def _bordes_parrafo(p, lados: dict[str, tuple[str, int, int]]) -> None:
    """lados = {"left": (color, grosor en 1/8 pt, separación en pt)}."""
    bdr = OxmlElement("w:pBdr")
    for lado in ("top", "left", "bottom", "right"):
        if lado in lados:
            color, grosor, espacio = lados[lado]
            bdr.append(_elemento(f"w:{lado}", val="single", sz=grosor, space=espacio, color=color))
    _insertar(p._p.get_or_add_pPr(), bdr, "w:pBdr", _ORDEN_PPR)


def _bordes_tabla(tabla, lados: dict[str, tuple[str, int]]) -> None:
    """lados = {"left": (color, grosor)}; los que falten quedan sin borde."""
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if lado in lados:
            color, grosor = lados[lado]
            bordes.append(_elemento(f"w:{lado}", val="single", sz=grosor, space=0, color=color))
        else:
            bordes.append(_elemento(f"w:{lado}", val="nil"))
    _insertar(tabla._tbl.tblPr, bordes, "w:tblBorders", _ORDEN_TBLPR)


def _margenes_celdas(tabla, arriba: int, lados: int) -> None:
    """Relleno interior de las celdas, en veinteavos de punto."""
    mar = OxmlElement("w:tblCellMar")
    for lado, valor in (("top", arriba), ("left", lados), ("bottom", arriba), ("right", lados)):
        mar.append(_elemento(f"w:{lado}", w=valor, type="dxa"))
    _insertar(tabla._tbl.tblPr, mar, "w:tblCellMar", _ORDEN_TBLPR)


def _sombrear_celda(celda, color: str) -> None:
    _insertar(celda._tc.get_or_add_tcPr(), _elemento("w:shd", val="clear", color="auto", fill=color),
              "w:shd", _ORDEN_TCPR)


def _sombrear_run(run, color: str) -> None:
    _insertar(run._r.get_or_add_rPr(), _elemento("w:shd", val="clear", color="auto", fill=color),
              "w:shd", _ORDEN_RPR)


def _fuente_fija(objeto_font, nombre: str) -> None:
    """Fija la fuente también para Asia oriental y quita los atajos al tema.

    Los estilos de la plantilla usan fuentes "del tema" (asciiTheme…), que
    tienen prioridad sobre font.name: si no se quitan, Word ignora el cambio.
    """
    objeto_font.name = nombre
    rfonts = objeto_font._element.get_or_add_rPr().find(qn("w:rFonts"))
    if rfonts is None:
        return
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    rfonts.set(qn("w:eastAsia"), nombre)
    rfonts.set(qn("w:cs"), nombre)


# ------------------------------------------------------------ matemáticas
_M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _m(nombre: str, **attrs):
    el = OxmlElement(f"m:{nombre}")
    for k, v in attrs.items():
        el.set(f"{_M}{k}", str(v))
    return el


def _vacio(el) -> bool:
    return el is None or (len(el) == 0 and not (el.text or "").strip())


def _reparar_omml(raiz) -> None:
    """Completa los elementos obligatorios que mathml2omml no emite.

    - <m:rad> sin <m:deg>: el esquema lo exige; sin él Word y LibreOffice
      muestran un índice vacío (√ se ve como ∜□).
    - <m:nary> con <m:sub>/<m:sup> vacío: hay que marcarlo como oculto o
      aparece un recuadro fantasma sobre el sumatorio.
    """
    for rad in raiz.iter(f"{_M}rad"):
        if rad.find(f"{_M}deg") is None:
            props = rad.find(f"{_M}radPr")
            if props is None:
                props = _m("radPr")
                rad.insert(0, props)
            props.append(_m("degHide", val="1"))
            rad.insert(list(rad).index(props) + 1, _m("deg"))

    for nary in raiz.iter(f"{_M}nary"):
        props = nary.find(f"{_M}naryPr")
        if props is None:
            props = _m("naryPr")
            nary.insert(0, props)
        for parte, marca in (("sub", "subHide"), ("sup", "supHide")):
            el = nary.find(f"{_M}{parte}")
            if el is None:
                el = _m(parte)
                nary.insert(list(nary).index(nary.find(f"{_M}e")), el)
            if _vacio(el) and props.find(f"{_M}{marca}") is None:
                # naryPr: chr, limLoc, grow, subHide, supHide, ctrlPr
                ctrl = props.find(f"{_M}ctrlPr")
                nuevo = _m(marca, val="1")
                if ctrl is not None:
                    ctrl.addprevious(nuevo)
                else:
                    props.append(nuevo)

    # Runs vacíos (de espacios como \, o \;): no aportan nada.
    for r in list(raiz.iter(f"{_M}r")):
        t = r.find(f"{_M}t")
        if t is not None and not (t.text or "") and r.getparent() is not None:
            r.getparent().remove(r)


def _omml(latex: str, bloque: bool):
    """LaTeX → elemento OMML (ecuación nativa de Word), o None si falla."""
    try:
        mml = latex2mathml.converter.convert(latex, display="block" if bloque else "inline")
        omml = mathml2omml.convert(mml)
    except Exception:  # noqa: BLE001 — cualquier fallo cae al texto plano
        return None
    decl = nsdecls("m", "w")
    try:
        if bloque:
            el = parse_xml(f"<m:oMathPara {decl}>{omml}</m:oMathPara>")
        else:
            el = parse_xml(omml.replace("<m:oMath>", f"<m:oMath {decl}>", 1))
    except Exception:  # noqa: BLE001
        return None
    _reparar_omml(el)
    return el


# ---------------------------------------------------------------- imágenes
def _cargar_imagen(src: str, carpeta: Path) -> bytes | None:
    """Devuelve los bytes de una imagen local o remota, o None."""
    if not src:
        return None
    try:
        if src.startswith(("http://", "https://")):
            peticion = urllib.request.Request(src, headers={"User-Agent": "lectorMD"})
            with urllib.request.urlopen(peticion, timeout=10) as r:  # noqa: S310 — solo http/https
                return r.read(25 * 1024 * 1024)
        if src.startswith("data:"):
            _, datos = src.split(",", 1)
            return base64.b64decode(datos)
        if src.startswith("file://"):
            ruta = Path(unquote(urlparse(src).path))
        else:
            ruta = (carpeta / unquote(src)).resolve()
        return ruta.read_bytes()
    except (OSError, ValueError):
        return None


def _normalizar_imagen(datos: bytes) -> tuple[bytes, int] | None:
    """Asegura un formato que Word entienda y devuelve (bytes, ancho en px).

    Word no admite SVG ni WebP: si Qt está disponible se convierten a PNG.
    """
    try:
        img = ImagenDocx.from_blob(datos)
        if img.content_type.split("/")[-1] in FORMATOS_DOCX:
            return datos, img.px_width
    except Exception:  # noqa: BLE001 — formato no reconocido por python-docx
        pass
    try:
        from PySide6.QtCore import QBuffer, QByteArray, QIODevice
        from PySide6.QtGui import QImage

        qimg = QImage.fromData(QByteArray(datos))
        if qimg.isNull():
            return None
        buf = QBuffer()
        buf.open(QIODevice.OpenModeFlag.WriteOnly)
        qimg.save(buf, "PNG")
        return bytes(buf.data()), qimg.width()
    except ImportError:
        return None


# --------------------------------------------------------------- exportador
class ExportadorWord:
    """Construye un .docx a partir de un documento ya analizado."""

    def __init__(self, analisis: renderer.Analisis, carpeta: Path,
                 diagramas: list[dict | None] | None = None):
        self.a = analisis
        self.carpeta = carpeta
        self.diagramas = list(diagramas or [])
        self.n_diagrama = 0
        self.doc = Document()
        self.contenedores: list = [self.doc]   # documento o celda de una caja
        self.ultimo_parrafo = None
        self.listas: list[dict] = []
        self.citas: list[str | None] = []   # None = cita normal, "aviso" = caja
        self.primer_parrafo_item = False
        self.estilo_pygments = get_style_by_name(renderer.ESTILO_CLARO)
        self._preparar()

    # ------------------------------------------------------------- estilos
    def _preparar(self) -> None:
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = MARGEN
        sec.top_margin = sec.bottom_margin = Cm(2.4)

        normal = self.doc.styles["Normal"]
        _fuente_fija(normal.font, FUENTE)
        normal.font.size = Pt(11)
        normal.font.color.rgb = TEXTO
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.2

        tamanos = {1: 20, 2: 15.5, 3: 13, 4: 11.5, 5: 11, 6: 11}
        for nivel, tam in tamanos.items():
            est = self.doc.styles[f"Heading {nivel}"]
            _fuente_fija(est.font, FUENTE)
            est.font.size = Pt(tam)
            est.font.bold = True
            est.font.italic = False
            est.font.color.rgb = ACENTO if nivel <= 2 else TEXTO
            est.paragraph_format.space_before = Pt(18 if nivel <= 2 else 12)
            est.paragraph_format.space_after = Pt(6)

        titulo = self.doc.styles["Title"]
        _fuente_fija(titulo.font, FUENTE)
        titulo.font.size = Pt(26)
        titulo.font.bold = True
        titulo.font.color.rgb = TEXTO
        # La plantilla dibuja una raya bajo el título; se quita.
        ppr = titulo.element.get_or_add_pPr()
        for b in ppr.findall(qn("w:pBdr")):
            ppr.remove(b)

        # Declarar la fuente de código como de paso fijo: si el archivo se abre
        # donde no está instalada, Word elige otra monoespaciada.
        tabla_fuentes = next((pt for pt in self.doc.part.package.iter_parts()
                              if str(pt.partname).endswith("fontTable.xml")), None)
        nombre = FUENTE_MONO.encode()
        if tabla_fuentes is not None and b'w:name="' + nombre + b'"' not in tabla_fuentes.blob:
            panose = PANOSE_MONO.get(FUENTE_MONO, "")
            entrada = (b'<w:font w:name="' + nombre + b'">'
                       + (b'<w:panose1 w:val="' + panose.encode() + b'"/>' if panose else b"")
                       + b'<w:charset w:val="00"/><w:family w:val="modern"/>'
                       b'<w:pitch w:val="fixed"/></w:font>')
            if tabla_fuentes.blob.count(b"</w:fonts>") == 1:
                tabla_fuentes._blob = tabla_fuentes.blob.replace(b"</w:fonts>", entrada + b"</w:fonts>")

        for nombre in ("Quote", "List Bullet", "List Number", "List Bullet 2",
                       "List Number 2", "List Bullet 3", "List Number 3"):
            if nombre in [s.name for s in self.doc.styles]:
                _fuente_fija(self.doc.styles[nombre].font, FUENTE)

    # -------------------------------------------------------- contenedores
    def _c(self):
        return self.contenedores[-1]

    def _agregar_parrafo(self, texto: str = "", style=None):
        p = self._c().add_paragraph(texto, style)
        self.ultimo_parrafo = p
        return p

    def _caja(self, fondo: str, bordes: dict):
        """Tabla de una celda: fondo y bordes que no se fusionan con vecinos.

        Word une párrafos contiguos con bordes idénticos, así que dos bloques
        de código seguidos acabarían en una sola caja si fueran párrafos.
        """
        tabla = self._c().add_table(rows=1, cols=1)
        tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
        _bordes_tabla(tabla, bordes)
        _margenes_celdas(tabla, 90, 150)
        if self.listas:
            sangria = _elemento("w:tblInd", w=int(Cm(1.27 * len(self.listas)).twips), type="dxa")
            _insertar(tabla._tbl.tblPr, sangria, "w:tblInd", _ORDEN_TBLPR)
        celda = tabla.cell(0, 0)
        _sombrear_celda(celda, fondo)
        return celda

    def _separador(self) -> None:
        """Párrafo mínimo tras una tabla.

        Evita que Word funda dos tablas seguidas y cumple la regla de que una
        celda debe terminar en párrafo.
        """
        p = self._agregar_parrafo()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(4)
        pf.line_spacing = Pt(6)

    # ------------------------------------------------------------ listas
    def _estilo_lista(self, tipo: str, nivel: int) -> str:
        base = "List Bullet" if tipo == "bullet" else "List Number"
        return base if nivel == 0 else f"{base} {min(nivel, 2) + 1}"

    def _numeracion_nueva(self, estilo: str) -> int | None:
        """Crea una numeración que empieza en 1 para una lista ordenada.

        Sin esto Word continúa la cuenta de la lista anterior (1, 2, 3… 4, 5).
        """
        try:
            numpr = self.doc.styles[estilo].element.pPr.numPr
            numeracion = self.doc.part.numbering_part.numbering_definitions._numbering
            abstracta = numeracion.num_having_numId(numpr.numId.val).abstractNumId.val
            num = numeracion.add_num(abstracta)
            num.add_lvlOverride(ilvl=0).add_startOverride(1)
            return num.numId
        except (AttributeError, KeyError):
            return None

    # --------------------------------------------------------- párrafos
    def _parrafo(self):
        """Crea el párrafo adecuado según el contexto (lista, cita, normal)."""
        if self.listas:
            lista = self.listas[-1]
            nivel = len(self.listas) - 1
            if lista["tareas"]:
                p = self._agregar_parrafo()
                p.paragraph_format.left_indent = Cm(0.6 + 0.9 * nivel)
                p.paragraph_format.space_after = Pt(2)
            elif self.primer_parrafo_item:
                p = self._agregar_parrafo(style=self._estilo_lista(lista["tipo"], nivel))
                if lista["tipo"] == "ordered" and lista["num_id"]:
                    numpr = p._p.get_or_add_pPr().get_or_add_numPr()
                    numpr.get_or_add_ilvl().val = 0
                    numpr.get_or_add_numId().val = lista["num_id"]
                p.paragraph_format.space_after = Pt(2)
            else:
                # Segundo párrafo dentro del mismo punto: sangrado, sin viñeta.
                p = self._agregar_parrafo()
                p.paragraph_format.left_indent = Cm(1.27 + 0.63 * nivel)
            self.primer_parrafo_item = False
        else:
            p = self._agregar_parrafo()

        if self.citas and self.citas[-1] is None:
            self._estilo_cita(p)
        return p

    def _estilo_cita(self, p) -> None:
        pf = p.paragraph_format
        normales = sum(1 for c in self.citas if c is None)
        pf.left_indent = (pf.left_indent or 0) + Cm(0.5 * normales)
        pf.space_after = Pt(4)
        _bordes_parrafo(p, {"left": ("C8CCD2", 14, 10)})
        p._cita = True  # marca para poner cursiva en los runs

    # ----------------------------------------------------------- en línea
    def _run(self, p, destino, texto: str, fmt: dict):
        r = p.add_run(texto)
        if fmt.get("b"):
            r.bold = True
        if fmt.get("i") or getattr(p, "_cita", False):
            r.italic = True
        if fmt.get("s"):
            r.font.strike = True
        if fmt.get("sup"):
            r.font.superscript = True
        if fmt.get("sub"):
            r.font.subscript = True
        if fmt.get("color"):
            r.font.color.rgb = fmt["color"]
        if fmt.get("enlace"):
            r.font.color.rgb = ACENTO
            r.font.underline = True
        if fmt.get("mono"):
            _fuente_fija(r.font, FUENTE_MONO)
            r.font.size = Pt(9.5)
            _sombrear_run(r, FONDO_CODIGO_LINEA)
        if destino is not None:
            destino.append(r._r)
        return r

    def _hipervinculo(self, p, url: str):
        r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        p._p.append(h)
        return h

    def _en_linea(self, p, hijos) -> None:
        fmt: dict = {}
        destino = None
        solo_imagen = [h.type for h in hijos if not (h.type == "text" and not h.content.strip())] == ["image"]

        for h in hijos:
            t = h.type
            if t == "text":
                if h.content:
                    self._run(p, destino, h.content, fmt)
            elif t == "softbreak":
                self._run(p, destino, " ", fmt)
            elif t == "hardbreak":
                self._run(p, destino, "", fmt).add_break()
            elif t in ("strong_open", "strong_close"):
                fmt["b"] = t.endswith("open")
            elif t in ("em_open", "em_close"):
                fmt["i"] = t.endswith("open")
            elif t in ("s_open", "s_close"):
                fmt["s"] = t.endswith("open")
            elif t == "code_inline":
                self._run(p, destino, h.content, {**fmt, "mono": True})
            elif t == "link_open":
                url = h.attrGet("href") or ""
                if url.startswith(("http://", "https://", "mailto:")):
                    destino = self._hipervinculo(p, url)
                fmt["enlace"] = True
            elif t == "link_close":
                destino = None
                fmt["enlace"] = False
            elif t == "image":
                self._imagen(p, h.attrGet("src") or "", h.content or "", solo_imagen)
            elif t == "math_inline":
                self._formula_en_linea(p, destino, h.content, h.meta.get("display", False))
            elif t == "html_inline":
                self._html_en_linea(p, destino, h.content, fmt)

    def _html_en_linea(self, p, destino, html: str, fmt: dict) -> None:
        etiqueta = html.strip().lower()
        if etiqueta.startswith("<input") and "checkbox" in etiqueta:
            marca = "☑ " if "checked" in etiqueta else "☐ "
            self._run(p, destino, marca, {"color": ACENTO if "checked" in etiqueta else SUAVE})
        elif re.match(r"<br\s*/?>", etiqueta):
            self._run(p, destino, "", fmt).add_break()
        elif etiqueta in ("<sup>", "</sup>"):
            fmt["sup"] = not etiqueta.startswith("</")
        elif etiqueta in ("<sub>", "</sub>"):
            fmt["sub"] = not etiqueta.startswith("</")
        elif etiqueta in ("<kbd>", "</kbd>", "<code>", "</code>"):
            fmt["mono"] = not etiqueta.startswith("</")
        elif etiqueta in ("<b>", "<strong>", "</b>", "</strong>"):
            fmt["b"] = not etiqueta.startswith("</")
        elif etiqueta in ("<i>", "<em>", "</i>", "</em>"):
            fmt["i"] = not etiqueta.startswith("</")

    def _formula_en_linea(self, p, destino, latex: str, display: bool) -> None:
        el = _omml(latex, bloque=False)
        if el is None:
            r = self._run(p, destino, latex, {"i": True})
            r.font.name = "Cambria Math"
            return
        (destino if destino is not None else p._p).append(el)

    def _imagen(self, p, src: str, alt: str, sola: bool) -> None:
        datos = _cargar_imagen(src, self.carpeta)
        norm = _normalizar_imagen(datos) if datos else None
        if not norm:
            self._run(p, None, f"[imagen: {alt or src}]", {"i": True, "color": SUAVE})
            return
        blob, ancho_px = norm
        ancho = min(Emu(ancho_px * EMU_POR_PX), ANCHO_UTIL)
        p.add_run().add_picture(io.BytesIO(blob), width=ancho)
        if sola:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ------------------------------------------------------------ bloques
    def _codigo(self, codigo: str, lenguaje: str) -> None:
        codigo = codigo.rstrip("\n").expandtabs(4)
        lexer = None
        if lenguaje:
            try:
                lexer = get_lexer_by_name(lenguaje)
            except ClassNotFound:
                lexer = None
        if lexer is None and len(codigo) > 40:
            try:
                lexer = guess_lexer(codigo)
            except ClassNotFound:
                lexer = None

        celda = self._caja(FONDO_CODIGO, {lado: (BORDE, 4) for lado in ("top", "left", "bottom", "right")})
        p = celda.paragraphs[0]
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0

        def escribir(texto: str, color=None, negrita=False, cursiva=False):
            partes = texto.split("\n")
            for n, parte in enumerate(partes):
                if n:
                    p.add_run().add_break()
                if parte:
                    r = p.add_run(parte)
                    _fuente_fija(r.font, FUENTE_MONO)
                    r.font.size = Pt(9)
                    if color:
                        r.font.color.rgb = RGBColor.from_string(color)
                    r.bold = negrita or None
                    r.italic = cursiva or None

        if lexer is None:
            escribir(codigo, color="24292F")
        else:
            for tipo, valor in lexer.get_tokens(codigo):
                estilo = self.estilo_pygments.style_for_token(tipo)
                escribir(valor, estilo.get("color") or "24292F", estilo.get("bold"), estilo.get("italic"))

            # Pygments añade un salto final; si quedó un <w:br/> colgando se quita.
            ultimo = p._p.findall(qn("w:r"))
            if ultimo and ultimo[-1].find(qn("w:br")) is not None and ultimo[-1].find(qn("w:t")) is None:
                p._p.remove(ultimo[-1])
        self._separador()

    def _diagrama(self, fuente: str) -> None:
        info = self.diagramas[self.n_diagrama] if self.n_diagrama < len(self.diagramas) else None
        self.n_diagrama += 1
        p = self._agregar_parrafo()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if info and info.get("png"):
            blob = base64.b64decode(info["png"].split(",", 1)[1])
            ancho = min(Emu(int(info.get("ancho", 600)) * EMU_POR_PX), ANCHO_UTIL)
            p.add_run().add_picture(io.BytesIO(blob), width=ancho)
        else:
            # Sin imagen (tipo de diagrama no exportable): se deja el código.
            p._p.getparent().remove(p._p)
            self._codigo(fuente, "text")

    def _formula_bloque(self, latex: str) -> None:
        p = self._agregar_parrafo()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        el = _omml(latex, bloque=True)
        if el is None:
            r = p.add_run(latex)
            r.font.name = "Cambria Math"
            r.italic = True
        else:
            p._p.append(el)

    def _regla(self) -> None:
        p = self._agregar_parrafo()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(12)
        _bordes_parrafo(p, {"bottom": ("C8CCD2", 6, 1)})

    def _tabla(self, tokens, inicio: int) -> int:
        """Construye una tabla; devuelve el índice tras table_close."""
        filas: list[list[tuple]] = []
        i = inicio + 1
        cabecera = False
        while tokens[i].type != "table_close":
            t = tokens[i]
            if t.type == "thead_open":
                cabecera = True
            elif t.type == "thead_close":
                cabecera = False
            elif t.type == "tr_open":
                filas.append([])
            elif t.type in ("th_open", "td_open"):
                estilo = t.attrGet("style") or ""
                m = re.search(r"text-align:\s*(left|center|right)", estilo)
                filas[-1].append((tokens[i + 1], m.group(1) if m else "left", cabecera))
            i += 1

        if not filas:
            return i + 1
        columnas = max(len(f) for f in filas)
        tabla = self._c().add_table(rows=len(filas), cols=columnas)
        tabla.style = self.doc.styles["Table Grid"]
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Bordes suaves en lugar del negro de la plantilla.
        _bordes_tabla(tabla, {lado: (BORDE, 4) for lado in
                              ("top", "left", "bottom", "right", "insideH", "insideV")})
        _margenes_celdas(tabla, 50, 110)

        alinear = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}
        for f, fila in enumerate(filas):
            for c in range(columnas):
                celda = tabla.cell(f, c)
                p = celda.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if c >= len(fila):
                    continue
                inline, alin, es_cab = fila[c]
                p.alignment = alinear[alin]
                self._en_linea(p, inline.children or [])
                if es_cab:
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = SUAVE
                    _sombrear_celda(celda, "F2F3F6")

        self._separador()
        return i + 1

    def _abrir_cita(self, tokens, i: int) -> None:
        """Detecta si la cita es un aviso [!NOTE] y pone su título."""
        color = None
        if i + 2 < len(tokens) and tokens[i + 1].type == "paragraph_open":
            inline = tokens[i + 2]
            m = re.match(r"^\s*\[!(\w+)\]\s*", inline.content)
            if m and m.group(1).lower() in AVISOS:
                etiqueta, color = AVISOS[m.group(1).lower()]
                self._quitar_marca_aviso(inline)
                celda = self._caja("F8F9FB", {"left": (color, 24)})
                p = celda.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                self._run(p, None, etiqueta, {"b": True, "color": RGBColor.from_string(color)})
                self.contenedores.append(celda)
                self.citas.append("aviso")
                return
        self.citas.append(None)

    @staticmethod
    def _quitar_marca_aviso(inline) -> None:
        hijos = list(inline.children or [])
        if hijos and hijos[0].type == "text":
            hijos[0].content = re.sub(r"^\s*\[!\w+\]\s*", "", hijos[0].content)
            if not hijos[0].content and len(hijos) > 1 and hijos[1].type in ("softbreak", "hardbreak"):
                hijos = hijos[2:]
            elif not hijos[0].content:
                hijos = hijos[1:]
        inline.children = hijos

    # ------------------------------------------------------------ recorrido
    def construir(self) -> Document:
        a = self.a
        if a.titulo:
            self._agregar_parrafo(a.titulo, style="Title")
            datos = [str(a.meta[k]) for k in ("autor", "author", "fecha", "date") if a.meta.get(k)]
            if datos:
                p = self._agregar_parrafo()
                self._run(p, None, "  ·  ".join(datos), {"color": SUAVE})
                p.paragraph_format.space_after = Pt(18)

        props = self.doc.core_properties
        props.title = a.titulo or ""
        props.author = str(a.meta.get("autor") or a.meta.get("author") or "")
        props.comments = "Exportado con lectorMD"

        tokens = a.tokens
        i = 0
        while i < len(tokens):
            t = tokens[i]
            tipo = t.type

            if tipo == "heading_open":
                nivel = int(t.tag[1])
                p = self._agregar_parrafo(style=f"Heading {nivel}")
                self._en_linea(p, tokens[i + 1].children or [])
                i += 3
                continue

            if tipo == "paragraph_open":
                p = self._parrafo()
                self._en_linea(p, tokens[i + 1].children or [])
                i += 3
                continue

            if tipo in ("bullet_list_open", "ordered_list_open"):
                clases = t.attrGet("class") or ""
                lista = {
                    "tipo": "bullet" if tipo.startswith("bullet") else "ordered",
                    "tareas": "lista-tareas" in clases,
                    "num_id": None,
                }
                if lista["tipo"] == "ordered":
                    lista["num_id"] = self._numeracion_nueva(self._estilo_lista("ordered", len(self.listas)))
                self.listas.append(lista)
            elif tipo in ("bullet_list_close", "ordered_list_close"):
                self.listas.pop()
                if not self.listas:
                    # Separación tras la lista completa.
                    if self.ultimo_parrafo is not None:
                        self.ultimo_parrafo.paragraph_format.space_after = Pt(8)
            elif tipo == "list_item_open":
                self.primer_parrafo_item = True
            elif tipo == "blockquote_open":
                self._abrir_cita(tokens, i)
            elif tipo == "blockquote_close":
                if self.citas.pop() == "aviso":
                    self.contenedores.pop()
                    self._separador()
            elif tipo in ("fence", "code_block"):
                lenguaje = (t.info or "").strip().split()[0] if t.info else ""
                if lenguaje.lower() == "mermaid":
                    self._diagrama(t.content)
                else:
                    self._codigo(t.content, lenguaje)
            elif tipo == "math_block":
                self._formula_bloque(t.content)
            elif tipo == "table_open":
                i = self._tabla(tokens, i)
                continue
            elif tipo == "hr":
                self._regla()
            elif tipo == "html_block":
                texto = re.sub(r"<!--.*?-->", "", t.content, flags=re.S)
                texto = re.sub(r"<[^>]+>", "", texto).strip()
                if texto:
                    self._parrafo().add_run(texto)
            i += 1

        return self.doc


def exportar_docx(texto_md: str, destino: Path, carpeta: Path,
                  diagramas: list[dict | None] | None = None) -> None:
    """Punto de entrada: Markdown → archivo .docx."""
    analisis = renderer.analizar(texto_md)
    ExportadorWord(analisis, carpeta, diagramas).construir().save(str(destino))
